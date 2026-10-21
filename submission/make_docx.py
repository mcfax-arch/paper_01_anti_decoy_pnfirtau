"""Post-process the pandoc-generated DOCX: A4 layout, fonts, tables, images, page numbers."""
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_margins(cell, top=20, bottom=20, left=40, right=40):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for tag, value in (('top', top), ('bottom', bottom), ('start', left), ('end', right)):
        node = OxmlElement(f'w:{tag}')
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_page_number_footer(section):
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' PAGE '
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def set_table_borders(table):
    """Three-line table: strong top/bottom, light inside horizontal, no verticals."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge, sz in (('top', 8), ('bottom', 8), ('insideH', 4)):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '000000')
        borders.append(el)
    for edge in ('left', 'right', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'none')
        el.set(qn('w:sz'), '0')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), 'auto')
        borders.append(el)
    tblPr.append(borders)


doc = Document('paper_en.docx')

# Page setup: A4, comfortable margins
for section in doc.sections:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    add_page_number_footer(section)

# Body font
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(10.5)
normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')

for level in (1, 2, 3):
    try:
        style = doc.styles[f'Heading {level}']
        style.font.name = 'Times New Roman'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.font.size = Pt({1: 15, 2: 12.5, 3: 11}[level])
    except KeyError:
        pass

title_style = doc.styles['Title']
title_style.font.name = 'Times New Roman'
title_style.font.size = Pt(18)
title_style.font.color.rgb = RGBColor(0, 0, 0)

# Tables: three-line-ish grid, compact font
for table in doc.tables:
    table.style = 'Table'
    table.autofit = True
    set_table_borders(table)
    ncols = len(table.columns)
    cell_font = Pt(7.5 if ncols <= 10 else (6.5 if ncols <= 14 else 6.0))
    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = cell_font
                    run.font.name = 'Times New Roman'

# Figures: scale to text width, keep aspect ratio
max_width = Inches(6.6)
for shape in doc.inline_shapes:
    if shape.width > max_width:
        ratio = max_width / shape.width
        shape.width = int(shape.width * ratio)
        shape.height = int(shape.height * ratio)

doc.save('paper_en.docx')

# Verification
doc2 = Document('paper_en.docx')
print('paragraphs', len(doc2.paragraphs))
print('tables', len(doc2.tables))
print('images', len(doc2.inline_shapes))
for i in range(min(4, len(doc2.paragraphs))):
    p = doc2.paragraphs[i]
    print(i, repr(p.style.name), repr(p.text[:90]))
